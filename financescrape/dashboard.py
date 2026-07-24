"""Build the Excel dashboard workbook from extracted figures and matched filings."""

from __future__ import annotations

from datetime import date
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .filings import Filing
from .financials import Figure

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14)
_SUBTLE = Font(color="808080")
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CURRENCY_FMT = "#,##0"
_RATIO_FMT = "0.00"


def _style_header_row(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER
        cell.alignment = Alignment(vertical="center")


def _autosize(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _is_currency(fig: Figure) -> bool:
    return fig.unit == "USD"


def _period_label(fig: Figure) -> str:
    fp = fig.fiscal_period or ""
    fy = fig.fiscal_year
    if fp == "FY":
        return f"FY{fy}" if fy else "FY"
    return f"{fp} {fy}" if fy else fp


def _build_summary(ws, company, ticker, cik, cutoff_date, filings) -> None:
    ws.title = "Summary"
    ws["A1"] = "SEC EDGAR Dashboard"
    ws["A1"].font = _TITLE_FONT
    rows = [
        ("Company", company),
        ("Ticker", ticker),
        ("CIK", cik),
        ("As of (cutoff)", cutoff_date.isoformat()),
    ]
    r = 3
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Matched filings (latest on or before cutoff)").font = Font(bold=True)
    r += 1
    headers = ["Form", "Filing Date", "Report Date", "Accession", "Document URL"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=r, column=c, value=h)
    _style_header_row(ws, r, len(headers))
    r += 1
    for form, filing in filings.items():
        if filing is None:
            ws.cell(row=r, column=1, value=form)
            ws.cell(row=r, column=2, value="(none before cutoff)").font = _SUBTLE
        else:
            ws.cell(row=r, column=1, value=filing.form)
            ws.cell(row=r, column=2, value=filing.filing_date.isoformat())
            ws.cell(row=r, column=3,
                    value=filing.report_date.isoformat() if filing.report_date else "")
            ws.cell(row=r, column=4, value=filing.accession_number)
            ws.cell(row=r, column=5, value=filing.document_url)
        r += 1

    _autosize(ws, {1: 22, 2: 24, 3: 16, 4: 24, 5: 60})


def _build_key_figures(ws, figures: list[Figure]) -> None:
    headers = ["Metric", "Period", "Value", "Unit", "Source Form"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    _style_header_row(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    r = 2
    for fig in figures:
        ws.cell(row=r, column=1, value=fig.label)
        ws.cell(row=r, column=2, value=_period_label(fig))
        val_cell = ws.cell(row=r, column=3, value=fig.value)
        if _is_currency(fig):
            val_cell.number_format = _CURRENCY_FMT
        else:
            val_cell.number_format = _RATIO_FMT
        ws.cell(row=r, column=4, value=fig.unit)
        ws.cell(row=r, column=5, value=fig.form)
        r += 1

    if not figures:
        ws.cell(row=2, column=1, value="No standardized XBRL figures found for this period.").font = _SUBTLE

    _autosize(ws, {1: 26, 2: 12, 3: 20, 4: 12, 5: 12})
    _append_ratios(ws, figures, start_row=r + 1)


def _annual_map(figures: list[Figure]) -> dict[str, float]:
    """Latest annual value per label, for ratio calculations."""
    out: dict[str, float] = {}
    for fig in figures:
        if fig.fiscal_period == "FY":
            out.setdefault(fig.label, fig.value)
    return out


def _append_ratios(ws, figures: list[Figure], start_row: int) -> None:
    vals = _annual_map(figures)

    def ratio(numer: str, denom: str):
        n, d = vals.get(numer), vals.get(denom)
        if n is None or d in (None, 0):
            return None
        return n / d

    computed = [
        ("Current Ratio", ratio("Current Assets", "Current Liabilities")),
        ("Net Margin", ratio("Net Income", "Revenue")),
        ("Gross Margin", ratio("Gross Profit", "Revenue")),
        ("Return on Equity", ratio("Net Income", "Stockholders' Equity")),
        ("Debt to Equity", ratio("Total Liabilities", "Stockholders' Equity")),
    ]
    computed = [(label, v) for label, v in computed if v is not None]
    if not computed:
        return

    r = start_row
    ws.cell(row=r, column=1, value="Computed Ratios (latest FY)").font = Font(bold=True)
    r += 1
    for c, h in enumerate(["Ratio", "Value"], start=1):
        ws.cell(row=r, column=c, value=h)
    _style_header_row(ws, r, 2)
    r += 1
    for label, value in computed:
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=round(value, 4))
        cell.number_format = _RATIO_FMT
        r += 1


def _build_raw_facts(ws, figures: list[Figure]) -> None:
    headers = ["Label", "Concept", "Period End", "Fiscal Period",
               "Fiscal Year", "Value", "Unit", "Form", "Accession"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    _style_header_row(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    r = 2
    for fig in figures:
        ws.cell(row=r, column=1, value=fig.label)
        ws.cell(row=r, column=2, value=fig.concept)
        ws.cell(row=r, column=3, value=fig.period_end.isoformat())
        ws.cell(row=r, column=4, value=fig.fiscal_period)
        ws.cell(row=r, column=5, value=fig.fiscal_year)
        ws.cell(row=r, column=6, value=fig.value)
        ws.cell(row=r, column=7, value=fig.unit)
        ws.cell(row=r, column=8, value=fig.form)
        ws.cell(row=r, column=9, value=fig.accession)
        r += 1

    _autosize(ws, {1: 24, 2: 40, 3: 14, 4: 12, 5: 12, 6: 20, 7: 12, 8: 10, 9: 24})


def build_workbook(
    out_path: str,
    company: str,
    ticker: str,
    cik: str,
    cutoff_date: date,
    figures: list[Figure],
    filings: dict[str, Optional[Filing]],
) -> str:
    """Write the dashboard workbook to ``out_path`` and return the path."""
    wb = Workbook()
    _build_summary(wb.active, company, ticker, cik, cutoff_date, filings)
    _build_key_figures(wb.create_sheet("Key Figures"), figures)
    _build_raw_facts(wb.create_sheet("Raw Facts"), figures)
    wb.save(out_path)
    return out_path
